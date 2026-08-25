from openpyxl import load_workbook

from annotate.export_second_annotator import build_subset, write_subset_xlsx, LABEL_CHOICES


def test_build_subset_produces_target_total():
    rows = [
        {"url": f"https://a.com/{i}/", "outlet": "a", "source_style": "credible", "headline_text": f"H{i}"}
        for i in range(50)
    ] + [
        {"url": f"https://b.com/{i}/", "outlet": "b", "source_style": "clickbait_heavy", "headline_text": f"H{i}"}
        for i in range(50)
    ]
    subset_rows, mapping_rows = build_subset(rows, target_total=10, seed=1)
    assert len(subset_rows) == 10
    assert len(mapping_rows) == 10


def test_build_subset_hides_url_and_outlet_from_subset_rows():
    rows = [
        {"url": "https://ghpage.com/1/", "outlet": "ghpage", "source_style": "clickbait_heavy", "headline_text": "H1"},
        {"url": "https://graphic.com.gh/2/", "outlet": "graphic", "source_style": "credible", "headline_text": "H2"},
    ]
    subset_rows, mapping_rows = build_subset(rows, target_total=2, seed=1)
    for row in subset_rows:
        assert set(row.keys()) == {"id", "headline_text", "label"}


def test_build_subset_mapping_allows_merge_back():
    rows = [{"url": "https://a.com/1/", "outlet": "a", "source_style": "credible", "headline_text": "H1"}]
    subset_rows, mapping_rows = build_subset(rows, target_total=1, seed=1)
    assert mapping_rows[0]["id"] == subset_rows[0]["id"]
    assert mapping_rows[0]["url"] == "https://a.com/1/"


def test_build_subset_is_reproducible_with_same_seed():
    rows = [
        {"url": f"https://a.com/{i}/", "outlet": "a", "source_style": "credible", "headline_text": f"H{i}"}
        for i in range(20)
    ]
    subset1, _ = build_subset(rows, target_total=5, seed=7)
    subset2, _ = build_subset(rows, target_total=5, seed=7)
    assert subset1 == subset2


def test_write_subset_xlsx_contains_headers_and_rows(tmp_path):
    subset_rows = [
        {"id": "H001", "headline_text": "First headline", "label": ""},
        {"id": "H002", "headline_text": "Second headline", "label": ""},
    ]
    path = str(tmp_path / "subset.xlsx")
    write_subset_xlsx(subset_rows, path)

    wb = load_workbook(path)
    ws = wb.active
    assert [c.value for c in ws[1]] == ["id", "headline_text", "label"]
    assert [c.value for c in ws[2]] == ["H001", "First headline", None]
    assert [c.value for c in ws[3]] == ["H002", "Second headline", None]


def test_write_subset_xlsx_adds_label_dropdown_on_every_data_row(tmp_path):
    subset_rows = [
        {"id": "H001", "headline_text": "First headline", "label": ""},
        {"id": "H002", "headline_text": "Second headline", "label": ""},
    ]
    path = str(tmp_path / "subset.xlsx")
    write_subset_xlsx(subset_rows, path)

    wb = load_workbook(path)
    ws = wb.active
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    dv = validations[0]
    assert dv.type == "list"
    assert dv.formula1 == '"' + ",".join(LABEL_CHOICES) + '"'
    assert "C2:C3" in str(dv.sqref)
