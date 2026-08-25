import csv
import os

from openpyxl import load_workbook


def read_csv_rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_xlsx_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    return [
        dict(zip(header, row))
        for row in rows_iter
        if any(cell is not None for cell in row)
    ]


def read_table_rows(path):
    if path.lower().endswith(".xlsx"):
        return read_xlsx_rows(path)
    return read_csv_rows(path)


def write_csv_rows(rows, path, fieldnames):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def append_csv_row(row, path, fieldnames):
    file_exists = os.path.exists(path)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if not file_exists:
            writer.writeheader()
        writer.writerow(row)
